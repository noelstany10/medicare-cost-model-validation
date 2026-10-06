"""Excel validation workbook: live formulas, conditional formatting, charts and a pivot-ready data table.

Validators/actuaries can audit every derived number in Excel: predictive ratios, PSI, segment breaches and the
SUMIFS cross-tab are formulas over the raw sums, not pasted values.
"""
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from . import config as C
from . import models as M
from .validation import RACE

HDR = PatternFill("solid", fgColor="1F3864")
HFONT = Font(bold=True, color="FFFFFF")
INPUT = PatternFill("solid", fgColor="EAF1FB")
GOOD, AMBER, BAD = (PatternFill("solid", fgColor=c) for c in ("D8F0D8", "FFF1C9", "F8D3D3"))
THIN = Border(bottom=Side(style="thin", color="D9D9D9"))


def _header(ws, row, cols, start_col=1):
    for j, c in enumerate(cols):
        cell = ws.cell(row=row, column=start_col + j, value=c)
        cell.fill, cell.font = HDR, HFONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=14, color="1F3864")
    if sub:
        ws["A2"] = sub
        ws["A2"].font = Font(italic=True, color="595959")


def _write_df(ws, df, row, col=1, fmt=None):
    _header(ws, row, list(df.columns), col)
    for i, r in enumerate(df.itertuples(index=False), 1):
        for j, v in enumerate(r):
            if isinstance(v, (np.floating, float)) and (np.isnan(v) or np.isinf(v)):
                v = None
            elif isinstance(v, np.generic):
                v = v.item()
            c = ws.cell(row=row + i, column=col + j, value=v)
            c.border = THIN
            if fmt and df.columns[j] in fmt:
                c.number_format = fmt[df.columns[j]]
    return row + len(df)


def build(R, dq, findings, rating, oot, meta):
    wb = Workbook()

    # ---------------- Scored members (pivot-ready) - built first so formulas can reference it
    sc = pd.DataFrame({
        "bene_id": oot["bene_id"].values,
        "age_band": M.age_band(oot["age"]).astype(str).values,
        "sex": np.where(oot["is_female"] == 1, "Female", "Male"),
        "race": oot["race_cd"].map(RACE).fillna("Unknown").values,
        "esrd": np.where(oot["esrd"] == 1, "Y", "N"),
        "n_chronic": oot["n_chronic"].astype(int).values,
        "prior_year_paid": oot["paid_total"].values,
        "predicted_cost": np.round(R["preds"]["oot"], 2),
        "actual_cost": oot[C.TARGET_COST].values,
        "hcc_probability": np.round(R["preds"]["oot_clf"], 4),
        "hcc_actual": oot[C.TARGET_HCC].values.astype(int),
    })
    sc["risk_decile"] = pd.qcut(sc.predicted_cost.rank(method="first"), 10, labels=range(1, 11)).astype(int)
    n = len(sc)

    ws = wb.active
    ws.title = "Summary"
    # placeholders for sheet order; filled below
    sheets = {name: wb.create_sheet(name) for name in
              ["Findings", "Data_Quality", "Model_Comparison", "Decile_Backtest", "PSI_Stability",
               "Segment_Calibration", "Sensitivity", "Pivot_Analysis", "Scored_Members"]}

    s = sheets["Scored_Members"]
    _write_df(s, sc, 1, fmt={"prior_year_paid": "#,##0", "predicted_cost": "#,##0", "actual_cost": "#,##0",
                             "hcc_probability": "0.0%"})
    tab = Table(displayName="tblScored", ref=f"A1:{get_column_letter(sc.shape[1])}{n + 1}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    s.add_table(tab)
    s.freeze_panes = "A2"
    _widths(s, [19, 10, 9, 10, 6, 10, 15, 15, 13, 15, 11, 11])
    rng = lambda col: f"Scored_Members!${col}$2:${col}${n + 1}"
    AGE, SEX, PRED, ACT, PROB, HCC, DEC = rng("B"), rng("C"), rng("H"), rng("I"), rng("J"), rng("K"), rng("L")

    # ---------------- Summary
    _title(ws, "Medicare Prospective Cost Model - Validation Workbook", f"Data: {meta['data_source']}  |  Run: {meta['run_date']}")
    rc = R["reg_comparison"].set_index(["model", "sample"])
    o = rc.loc[(M.CHAMPION_REG, "OOT 2009->2010")]
    cc = R["clf_comparison"].set_index(["model", "sample"]).loc[(M.CHAMPION_CLF, "OOT 2009->2010")]
    ws["A4"], ws["B4"] = "Overall validation rating", rating
    ws["A4"].font = ws["B4"].font = Font(bold=True)
    rows = [
        ("Out-of-time members scored", f"=COUNT({PRED})", "#,##0"),
        ("Total predicted 2010 cost ($)", f"=SUM({PRED})", "#,##0"),
        ("Total actual 2010 cost ($)", f"=SUM({ACT})", "#,##0"),
        ("Predictive ratio (pred / actual)", "=B7/B8", "0.000"),
        ("R-squared (OOT)", f"=RSQ({ACT},{PRED})", "0.000"),
        ("Cumming's Prediction Measure (OOT)", float(o.CPM), "0.000"),
        ("Mean absolute error ($)", float(o.MAE), "#,##0"),
        ("High-cost claimant AUC (OOT)", float(cc.AUC), "0.000"),
        ("High-cost lift, top decile", float(cc.Lift_top10), "0.00\"x\""),
        ("Score PSI (dev vs OOT)", "=PSI_Stability!H14", "0.000"),
        ("High findings", '=COUNTIF(Findings!C:C,"High")', "0"),
        ("Medium findings", '=COUNTIF(Findings!C:C,"Medium")', "0"),
        ("Low findings", '=COUNTIF(Findings!C:C,"Low")', "0"),
        ("Data-quality checks failed", '=COUNTIF(Data_Quality!F:F,"FAIL")', "0"),
    ]
    _header(ws, 5, ["Key metric", "Value"])
    for i, (k, v, f) in enumerate(rows, 6):
        ws.cell(row=i, column=1, value=k).border = THIN
        c = ws.cell(row=i, column=2, value=v); c.number_format = f; c.border = THIN
    ws.conditional_formatting.add("B9", CellIsRule(operator="between", formula=["0.95", "1.05"], fill=GOOD))
    ws.conditional_formatting.add("B9", CellIsRule(operator="notBetween", formula=["0.95", "1.05"], fill=AMBER))
    ws.conditional_formatting.add("B15", CellIsRule(operator="lessThan", formula=["0.1"], fill=GOOD))
    ws.conditional_formatting.add("B15", CellIsRule(operator="greaterThanOrEqual", formula=["0.1"], fill=AMBER))
    ws["A22"] = "Sheets: Findings | Data_Quality | Model_Comparison | Decile_Backtest | PSI_Stability | Segment_Calibration | Sensitivity | Pivot_Analysis | Scored_Members (Excel Table 'tblScored' - Insert > PivotTable)"
    ws["A22"].font = Font(italic=True, color="595959")
    _widths(ws, [38, 70])

    # ---------------- Findings
    f = sheets["Findings"]
    _title(f, "Validation findings & recommendations")
    end = _write_df(f, findings[["id", "area", "severity", "finding", "evidence", "recommendation"]], 3)
    f.conditional_formatting.add(f"C4:C{end}", CellIsRule(operator="equal", formula=['"High"'], fill=BAD))
    f.conditional_formatting.add(f"C4:C{end}", CellIsRule(operator="equal", formula=['"Medium"'], fill=AMBER))
    f.conditional_formatting.add(f"C4:C{end}", CellIsRule(operator="equal", formula=['"Low"'], fill=GOOD))
    _header(f, 3, ["ID", "Area", "Severity", "Finding", "Evidence", "Recommendation", "Owner", "Due date", "Status"])
    dv = DataValidation(type="list", formula1='"Open,In progress,Remediated,Accepted risk"', allow_blank=True)
    f.add_data_validation(dv)
    for r in range(4, end + 1):
        f.cell(row=r, column=9, value="Open").fill = INPUT
        dv.add(f.cell(row=r, column=9))
        for c in range(4, 7):
            f.cell(row=r, column=c).alignment = Alignment(wrap_text=True, vertical="top")
    _widths(f, [7, 18, 10, 40, 45, 50, 14, 12, 14])

    # ---------------- Data quality (status recomputed by formula from observed vs threshold)
    d = sheets["Data_Quality"]
    _title(d, "Data quality testing", "Status is a live formula: max-rule -> PASS if observed <= threshold, WARN if <= 5x, else FAIL")
    dqx = dq[["check_id", "category", "description", "observed", "threshold", "status", "rule"]].copy()
    end = _write_df(d, dqx, 3)
    for r in range(4, end + 1):
        d.cell(row=r, column=6, value=f'=IF(G{r}="info","INFO",IF(D{r}<=E{r},"PASS",IF(D{r}<=E{r}*5,"WARN","FAIL")))')
        d.cell(row=r, column=5).fill = INPUT
    for txt, fill in (("PASS", GOOD), ("WARN", AMBER), ("FAIL", BAD)):
        d.conditional_formatting.add(f"F4:F{end}", CellIsRule(operator="equal", formula=[f'"{txt}"'], fill=fill))
    _widths(d, [9, 15, 60, 14, 11, 9, 7])

    # ---------------- Model comparison
    m = sheets["Model_Comparison"]
    _title(m, "Benchmarking: cost models and high-cost claimant classifiers")
    end = _write_df(m, R["reg_comparison"].round(4), 3, fmt={"MAE": "#,##0", "RMSE": "#,##0", "Mean_actual": "#,##0", "Mean_predicted": "#,##0"})
    _write_df(m, R["clf_comparison"].round(4), end + 3)
    _widths(m, [32, 16] + [12] * 12)

    # ---------------- Decile backtest with formulas + chart
    b = sheets["Decile_Backtest"]
    _title(b, "Out-of-time decile backtest (2009 features -> 2010 actual)", "Sums are computed live from Scored_Members with SUMIFS")
    _header(b, 3, ["Decile", "Members", "Predicted $", "Actual $", "Mean predicted", "Mean actual", "Predictive ratio", "Cum. % actual cost", "HCC rate"])
    for i in range(1, 11):
        r = 3 + i
        b.cell(row=r, column=1, value=i)
        b.cell(row=r, column=2, value=f"=COUNTIFS({DEC},A{r})")
        b.cell(row=r, column=3, value=f"=SUMIFS({PRED},{DEC},A{r})").number_format = "#,##0"
        b.cell(row=r, column=4, value=f"=SUMIFS({ACT},{DEC},A{r})").number_format = "#,##0"
        b.cell(row=r, column=5, value=f"=C{r}/B{r}").number_format = "#,##0"
        b.cell(row=r, column=6, value=f"=D{r}/B{r}").number_format = "#,##0"
        b.cell(row=r, column=7, value=f"=C{r}/D{r}").number_format = "0.000"
        b.cell(row=r, column=8, value=f"=SUM($D$4:D{r})/SUM($D$4:$D$13)").number_format = "0.0%"
        b.cell(row=r, column=9, value=f"=AVERAGEIFS({HCC},{DEC},A{r})").number_format = "0.0%"
    b["A14"], b["B14"], b["C14"], b["D14"], b["G14"] = "Total", "=SUM(B4:B13)", "=SUM(C4:C13)", "=SUM(D4:D13)", "=C14/D14"
    b["C14"].number_format = b["D14"].number_format = "#,##0"; b["G14"].number_format = "0.000"
    b.conditional_formatting.add("G4:G14", CellIsRule(operator="notBetween", formula=["0.85", "1.15"], fill=AMBER))
    ch = BarChart(); ch.type = "col"; ch.title = "Mean predicted vs actual cost by risk decile"
    ch.add_data(Reference(b, min_col=5, max_col=6, min_row=3, max_row=13), titles_from_data=True)
    ch.set_categories(Reference(b, min_col=1, min_row=4, max_row=13))
    ch.y_axis.title = "$ per member"; ch.x_axis.title = "Risk decile"; ch.height, ch.width = 8, 16
    ch.series[0].graphicalProperties.solidFill = "2A78D6"; ch.series[1].graphicalProperties.solidFill = "EB6834"
    b.add_chart(ch, "K3")
    _widths(b, [8, 10, 14, 14, 14, 14, 14, 16, 10])

    # ---------------- PSI with formulas
    p = sheets["PSI_Stability"]
    _title(p, "Population Stability Index - champion score", "PSI contribution = (Actual% - Expected%) x LN(Actual% / Expected%)")
    t = R["psi_score_table"].copy()
    _header(p, 3, ["Bin", "Lower", "Upper", "Expected % (dev)", "Actual % (OOT)", "Difference", "LN ratio", "PSI contribution"])
    for i, row in enumerate(t.itertuples(index=False), 4):
        p.cell(row=i, column=1, value=int(row.bin))
        p.cell(row=i, column=2, value=None if np.isinf(row.lower) else float(row.lower)).number_format = "#,##0"
        p.cell(row=i, column=3, value=None if np.isinf(row.upper) else float(row.upper)).number_format = "#,##0"
        p.cell(row=i, column=4, value=float(row.expected_pct)).number_format = "0.00%"
        p.cell(row=i, column=5, value=float(row.actual_pct)).number_format = "0.00%"
        p.cell(row=i, column=6, value=f"=E{i}-D{i}").number_format = "0.00%"
        p.cell(row=i, column=7, value=f"=LN(E{i}/D{i})").number_format = "0.0000"
        p.cell(row=i, column=8, value=f"=F{i}*G{i}").number_format = "0.0000"
    last = 3 + len(t)
    p["G14"], p["H14"] = "Total PSI", f"=SUM(H4:H{last})"
    p["G15"], p["H15"] = "Assessment", '=IF(H14<0.1,"Stable",IF(H14<0.25,"Monitor","Significant shift"))'
    p["G14"].font = p["G15"].font = Font(bold=True); p["H14"].number_format = "0.0000"
    p.conditional_formatting.add("H15", FormulaRule(formula=['H14<0.1'], fill=GOOD))
    p.conditional_formatting.add("H15", FormulaRule(formula=['H14>=0.1'], fill=AMBER))
    _write_df(p, R["csi"].head(15).round(4), 18)
    p.cell(row=17, column=1, value="Characteristic Stability Index (top 15 features)").font = Font(bold=True)
    _widths(p, [24, 12, 12, 16, 15, 12, 12, 16])

    # ---------------- Segment calibration with breach formula
    g = sheets["Segment_Calibration"]
    _title(g, "Segment calibration (OOT)", "Breach if n >= 500 and predictive ratio outside 0.85 - 1.15 (tolerances editable in blue cells)")
    g["J3"], g["K3"], g["J4"], g["K4"], g["J5"], g["K5"] = "Lower tolerance", 0.85, "Upper tolerance", 1.15, "Min n", 500
    for c in ("K3", "K4", "K5"):
        g[c].fill = INPUT
    sg = R["subgroups_oot"][["dimension", "level", "n", "actual_mean", "pred_mean"]].copy()
    end = _write_df(g, sg, 3, fmt={"actual_mean": "#,##0", "pred_mean": "#,##0"})
    _header(g, 3, ["Dimension", "Level", "Members", "Actual mean $", "Predicted mean $", "Predictive ratio", "Status"])
    for r in range(4, end + 1):
        g.cell(row=r, column=6, value=f"=E{r}/D{r}").number_format = "0.000"
        g.cell(row=r, column=7, value=f'=IF(AND(C{r}>=$K$5,OR(F{r}<$K$3,F{r}>$K$4)),"BREACH","OK")')
    g.conditional_formatting.add(f"G4:G{end}", CellIsRule(operator="equal", formula=['"BREACH"'], fill=BAD))
    g.conditional_formatting.add(f"G4:G{end}", CellIsRule(operator="equal", formula=['"OK"'], fill=GOOD))
    _widths(g, [18, 12, 10, 14, 16, 15, 10, 2, 2, 16, 8])

    # ---------------- Sensitivity / robustness
    v = sheets["Sensitivity"]
    _title(v, "Sensitivity & robustness testing")
    end = _write_df(v, R["sensitivity"].round(4), 3, fmt={"mean_pred_change": "0.0%"})
    _write_df(v, R["robustness"].round(4), end + 3)
    _widths(v, [32, 16, 10, 10, 16, 10])

    # ---------------- Pivot analysis (dynamic cross-tab via SUMIFS / COUNTIFS)
    pv = sheets["Pivot_Analysis"]
    _title(pv, "Predictive ratio cross-tab: age band x sex (live SUMIFS over Scored_Members)",
           "Select a risk decile filter in B3 (blank = all). For ad-hoc slicing use Insert > PivotTable on table 'tblScored'.")
    pv["A3"], pv["B3"] = "Risk decile filter", None
    pv["B3"].fill = INPUT
    dv2 = DataValidation(type="list", formula1='"1,2,3,4,5,6,7,8,9,10"', allow_blank=True); pv.add_data_validation(dv2); dv2.add("B3")
    bands = ["<65", "65-69", "70-74", "75-79", "80-84", "85+"]
    _header(pv, 5, ["Age band", "Female PR", "Male PR", "All PR", "Members", "Actual $ / member"])
    for i, a in enumerate(bands, 6):
        pv.cell(row=i, column=1, value=a)
        dfilt = f',{DEC},IF($B$3="","<>",$B$3)'
        dsel = f'IF($B$3="",">0",$B$3)'
        for j, sx in ((2, "Female"), (3, "Male")):
            pv.cell(row=i, column=j, value=f'=IFERROR(SUMIFS({PRED},{AGE},$A{i},{SEX},"{sx}",{DEC},{dsel})/SUMIFS({ACT},{AGE},$A{i},{SEX},"{sx}",{DEC},{dsel}),"-")').number_format = "0.000"
        pv.cell(row=i, column=4, value=f'=IFERROR(SUMIFS({PRED},{AGE},$A{i},{DEC},{dsel})/SUMIFS({ACT},{AGE},$A{i},{DEC},{dsel}),"-")').number_format = "0.000"
        pv.cell(row=i, column=5, value=f'=COUNTIFS({AGE},$A{i},{DEC},{dsel})').number_format = "#,##0"
        pv.cell(row=i, column=6, value=f'=IFERROR(SUMIFS({ACT},{AGE},$A{i},{DEC},{dsel})/E{i},"-")').number_format = "#,##0"
    pv.conditional_formatting.add("B6:D11", CellIsRule(operator="notBetween", formula=["0.85", "1.15"], fill=AMBER))
    lc = LineChart(); lc.title = "Predictive ratio by age band"; lc.height, lc.width = 7, 14
    lc.add_data(Reference(pv, min_col=2, max_col=3, min_row=5, max_row=11), titles_from_data=True)
    lc.set_categories(Reference(pv, min_col=1, min_row=6, max_row=11))
    pv.add_chart(lc, "H5")
    _widths(pv, [18, 12, 12, 12, 12, 16])

    C.EXCEL_DIR.mkdir(parents=True, exist_ok=True)
    path = C.EXCEL_DIR / "Model_Validation_Workbook.xlsx"
    wb.save(path)
    return path
